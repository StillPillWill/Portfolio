from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("portfolio-build-brief.docx")

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1C2430"
MUTED = "5D6A7A"
LIGHT = "F2F4F7"
BLUE_GRAY = "E8EEF5"
LINE = "D5DCE5"
PLACEHOLDER = "6F7B8A"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths_dxa = [int(round(width * 1440)) for width in widths]
    total = sum(widths_dxa)

    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout"):
        existing = tbl_pr.first_child_found_in(f"w:{tag}")
        if existing is not None:
            tbl_pr.remove(existing)

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)

    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_paragraph(paragraph, before=0, after=6, line=1.25, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def add_cell_text(cell, text, size=10.5, color=INK, bold=False, italic=False, after=0):
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, after=after, line=1.15)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_placeholder(cell, text="[TYPE HERE]", lines=1):
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, after=0, line=1.15)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=PLACEHOLDER, italic=True)
    for _ in range(lines - 1):
        run = paragraph.add_run("\n[TYPE HERE]")
        set_run_font(run, size=10.5, color=PLACEHOLDER, italic=True)
    return paragraph


def add_label_value_table(doc, rows, widths=(1.75, 4.75)):
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, widths)
    for label, value, lines in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT)
        add_cell_text(cells[0], label.upper(), size=8.8, color=DARK_BLUE, bold=True)
        if value is None:
            add_placeholder(cells[1], lines=lines)
        else:
            add_cell_text(cells[1], value, size=10.2, color=PLACEHOLDER, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_checklist_table(doc, items):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, (6.5,))
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, BLUE_GRAY)
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, after=0, line=1.2)
    for index, item in enumerate(items):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(item)
        set_run_font(run, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, (6.5,))
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, BLUE_GRAY)
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, after=0, line=1.2)
    label_run = paragraph.add_run(label.upper() + "  ")
    set_run_font(label_run, size=8.8, color=BLUE, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_project_card(doc, group, number, defaults=None):
    defaults = defaults or {}
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, (1.65, 4.85))
    header = table.rows[0].cells
    header[0].merge(header[1])
    set_cell_shading(header[0], NAVY)
    paragraph = clear_cell(header[0])
    set_paragraph(paragraph, after=0, line=1.0, keep=True)
    run = paragraph.add_run(f"{group.upper()} PROJECT {number:02d}")
    set_run_font(run, size=9.2, color=WHITE, bold=True)

    fields = [
        ("Working name", defaults.get("name"), 1),
        ("Repo / URL", defaults.get("repo"), 1),
        ("Category / focus", None, 1),
        ("Why it matters", None, 2),
        ("What I actually built", None, 2),
        ("Proof / outcome", None, 2),
        ("Media to show", None, 2),
        ("Include?", "[ ] Yes   [ ] No   [ ] Unsure", 1),
    ]
    for label, value, lines in fields:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT)
        add_cell_text(cells[0], label.upper(), size=8.6, color=DARK_BLUE, bold=True)
        if value is None:
            add_placeholder(cells[1], lines=lines)
        else:
            add_cell_text(cells[1], value, size=10.2, color=PLACEHOLDER, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_section_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_paragraph(paragraph, keep=True)
    return paragraph


def add_body(doc, text, italic=False, color=INK, after=6):
    paragraph = doc.add_paragraph(style="Normal")
    set_paragraph(paragraph, after=after)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=color, italic=italic)
    return paragraph


def add_kicker(doc, text):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, after=6, line=1.0, keep=True)
    run = paragraph.add_run(text.upper())
    set_run_font(run, size=9, color=BLUE, bold=True)
    return paragraph


def add_title(doc, text):
    paragraph = doc.add_paragraph(style="Brief Title")
    set_paragraph(paragraph, after=5, line=1.0, keep=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=28, color=NAVY, bold=True)
    return paragraph


def add_subtitle(doc, text):
    paragraph = doc.add_paragraph(style="Brief Subtitle")
    set_paragraph(paragraph, after=18, line=1.15, keep=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=13, color=MUTED)
    return paragraph


def add_page_break(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)
    return paragraph


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True

    for name, size, color in (("Brief Title", 28, NAVY), ("Brief Subtitle", 13, MUTED)):
        style = styles.add_style(name, 1)
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    set_paragraph(header, after=0, line=1.0)
    header_run = header.add_run("STILLPILLWILL / PORTFOLIO BUILD BRIEF")
    set_run_font(header_run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    set_paragraph(footer, after=0, line=1.0)
    left = footer.add_run("Working brief - fill and return")
    set_run_font(left, size=8.5, color=MUTED)
    tab = footer.add_run("\tPage ")
    set_run_font(tab, size=8.5, color=MUTED)
    add_page_field(footer)


def build():
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Portfolio Build Brief"
    props.subject = "Fill-in brief for the next local portfolio implementation"
    props.author = "StillPillWill"
    props.keywords = "portfolio, website, project inventory, visual direction"

    add_kicker(doc, "STILLPILLWILL / LOCAL PORTFOLIO")
    add_title(doc, "Portfolio Build Brief")
    add_subtitle(doc, "A fill-in source of truth for the next localhost implementation.")
    add_callout(
        doc,
        "How to use",
        "Replace every [TYPE HERE] field. Edit the prefilled assumptions if they are wrong. If a field does not apply, write N/A instead of leaving me to guess.",
    )

    add_section_heading(doc, "1. Start here: decisions I should not guess", 1)
    add_body(doc, "These are the highest-impact facts. Keep the answers short and direct; they will become the site copy and structure.")
    add_label_value_table(
        doc,
        [
            ("Name shown on site", "William Nzive - confirm or replace", 1),
            ("Handle / GitHub", "StillPillWill - confirm or replace", 1),
            ("Primary role", None, 1),
            ("One-line positioning", None, 1),
            ("Primary contact link", None, 1),
            ("Resume / CV link", None, 1),
            ("The site should feel like", None, 1),
            ("Words or ideas to avoid", None, 1),
        ],
    )

    add_section_heading(doc, "Known constraints - confirm or edit", 2)
    add_checklist_table(
        doc,
        [
            "[x] Keep the site local during this iteration",
            "[x] Separate Computer Science and Mechanics work",
            "[x] Keep navigation spatial; no normal document scroll",
            "[x] Exclude vibe-coded or weak showcase projects",
            "[ ] Other non-negotiable: [TYPE HERE]",
        ],
    )

    add_section_heading(doc, "What success looks like", 2)
    add_label_value_table(
        doc,
        [
            ("A visitor should understand", None, 2),
            ("The strongest proof should be", None, 2),
            ("The action I want visitors to take", None, 1),
        ],
    )

    add_page_break(doc)
    add_kicker(doc, "IDENTITY / VOICE")
    add_section_heading(doc, "2. Identity and voice", 1)
    add_body(doc, "Write in your natural voice. These answers determine the headline, about section, project descriptions, and contact language.")
    add_label_value_table(
        doc,
        [
            ("Short bio", None, 3),
            ("Long bio / background", None, 5),
            ("Three words for my voice", None, 1),
            ("What I care about in my work", None, 3),
            ("What makes my approach different", None, 3),
            ("What I do not want to sound like", None, 2),
        ],
    )
    add_section_heading(doc, "Links and proof", 2)
    add_label_value_table(
        doc,
        [
            ("GitHub", "https://github.com/StillPillWill - confirm or replace", 1),
            ("LinkedIn", None, 1),
            ("Email", None, 1),
            ("Other link", None, 1),
            ("Resume filename / path", None, 1),
        ],
    )
    add_callout(doc, "Writing rule", "If a project was a quick experiment, an existing tool, or something you would not defend in conversation, mark it No in the project section. It should not appear in the final site.")

    add_page_break(doc)
    add_kicker(doc, "PROJECTS / COMPUTER SCIENCE")
    add_section_heading(doc, "3. Computer Science project inventory", 1)
    add_body(doc, "Only include work you would be comfortable explaining in an interview or serious conversation. Replace or remove the current assumptions below.")
    add_project_card(doc, "Computer Science", 1, {"name": "DaT Challenge - confirm or replace", "repo": "https://github.com/StillPillWill/Dat-Challenge - confirm or replace"})
    add_project_card(doc, "Computer Science", 2, {"name": "CSI PPS - confirm or replace", "repo": "https://github.com/StillPillWill/Research - confirm or replace"})
    add_project_card(doc, "Computer Science", 3)
    add_callout(doc, "Selection test", "If a project is mostly generated, mostly copied, or not materially yours, write No in Include? and explain why here: [TYPE HERE]")

    add_page_break(doc)
    add_kicker(doc, "PROJECTS / MECHANICS")
    add_section_heading(doc, "4. Mechanics project inventory", 1)
    add_body(doc, "Describe the real build, the constraint, and the proof. The strongest mechanical projects usually have a clear before, change, and result.")
    add_project_card(doc, "Mechanics", 1, {"name": "Vulcan - confirm or replace", "repo": "https://github.com/StillPillWill/Vulcan - confirm or replace"})
    add_project_card(doc, "Mechanics", 2, {"name": "Ender3-2 - confirm or replace", "repo": "https://github.com/StillPillWill/Ender3-2 - confirm or replace"})
    add_project_card(doc, "Mechanics", 3)
    add_callout(doc, "Selection test", "If the project is not important enough to earn screen space, write No in Include? and leave it out of the final implementation.")

    add_page_break(doc)
    add_kicker(doc, "MEDIA / PRESENTATION")
    add_section_heading(doc, "5. Media and presentation", 1)
    add_body(doc, "For each included project, tell me which asset deserves the most space. Paste a local path, a GitHub raw URL, or a note that the repository has no usable media.")
    media_table = doc.add_table(rows=1, cols=4)
    set_table_geometry(media_table, (1.25, 2.1, 1.05, 2.1))
    mark_header_row(media_table.rows[0])
    headers = ["PROJECT", "ASSET PATH / URL", "TYPE", "CAPTION / PRESENTATION NOTE"]
    for cell, text in zip(media_table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        add_cell_text(cell, text, size=8.3, color=WHITE, bold=True)
    for _ in range(6):
        cells = media_table.add_row().cells
        for index, cell in enumerate(cells):
            if index == 2:
                add_placeholder(cell, "image / video / repo", 1)
            else:
                add_placeholder(cell, lines=2 if index == 3 else 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_section_heading(doc, "Media decisions", 2)
    add_checklist_table(
        doc,
        [
            "[ ] Preserve full image aspect ratio, even if that leaves some negative space",
            "[ ] Crop images when the subject becomes more legible",
            "[ ] Lead with the strongest image; keep extra assets as selectable media",
            "[ ] Autoplay project video when visible",
            "[ ] Never show low-resolution, duplicate, or contextless images",
            "[ ] Other media rule: [TYPE HERE]",
        ],
    )
    add_label_value_table(
        doc,
        [
            ("Images I definitely want", None, 2),
            ("Images I definitely do not want", None, 2),
            ("Preferred captions / credits", None, 2),
        ],
    )

    add_page_break(doc)
    add_kicker(doc, "EXPERIENCE / NAVIGATION")
    add_section_heading(doc, "6. How the site should behave", 1)
    add_body(doc, "Be specific about the experience. These answers determine the spatial field order, labels, transitions, and mobile behavior.")
    add_label_value_table(
        doc,
        [
            ("Navigation order", "Start -> CS -> Mechanics -> Method -> Signal - confirm or replace", 1),
            ("How movement should feel", None, 2),
            ("Desktop interaction", None, 2),
            ("Mobile interaction", None, 2),
            ("What should happen on project selection", None, 2),
            ("What should happen on media selection", None, 2),
            ("Accessibility / keyboard needs", None, 2),
            ("Performance concerns", None, 2),
        ],
    )
    add_section_heading(doc, "Reference sites", 2)
    add_label_value_table(
        doc,
        [
            ("Site 1 - what to borrow", None, 2),
            ("Site 2 - what to borrow", None, 2),
            ("Site 3 - what to avoid", None, 2),
        ],
    )

    add_page_break(doc)
    add_kicker(doc, "VISUAL DIRECTION")
    add_section_heading(doc, "7. What the site should look and feel like", 1)
    add_label_value_table(
        doc,
        [
            ("Palette in plain language", "Midnight is a starting point - define what that means to you: [TYPE HERE]", 2),
            ("Background color / mood", None, 1),
            ("Accent color(s)", None, 1),
            ("Typography preference", None, 1),
            ("Density preference", None, 1),
            ("Amount of motion", None, 1),
        ],
    )
    add_section_heading(doc, "Visual yes / no list", 2)
    add_checklist_table(
        doc,
        [
            "[ ] Clean and minimal",
            "[ ] Editorial / typographic",
            "[ ] Technical / precise",
            "[ ] Atmospheric / cinematic",
            "[ ] Warm / human",
            "[ ] High contrast",
            "[ ] Avoid neon cyan",
            "[ ] Avoid visible grid backgrounds",
            "[ ] Avoid large unused empty areas",
            "[ ] Avoid generic dashboard styling",
            "[ ] Other visual no-go: [TYPE HERE]",
        ],
    )
    add_section_heading(doc, "Final approval gate", 2)
    add_label_value_table(
        doc,
        [
            ("The homepage is successful when", None, 2),
            ("The project pages are successful when", None, 2),
            ("Before implementation, ask me about", None, 2),
            ("Anything else I should know", None, 3),
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
